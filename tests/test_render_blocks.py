from docx.shared import Pt

from generator import ReportGenerator
from models import Block, BlockKind


def test_render_two_blocks():
    """2 абзаца в документе."""
    blocks = [
        Block(kind=BlockKind.TITLE, text="Заголовок"),
        Block(kind=BlockKind.BODY, text="Текст"),
    ]

    gen = ReportGenerator()
    gen.render_blocks(blocks)
    assert len(gen.doc.paragraphs) == 2
    assert gen.doc.paragraphs[0].text == "Заголовок"
    assert gen.doc.paragraphs[1].text == "Текст"

def test_render_empty_after_code():
    """После CODE идет пустая строка."""
    blocks = [
        Block(kind=BlockKind.CODE, text="code"),
        Block(kind=BlockKind.SECTION, text="Выводы"),
    ]
    gen = ReportGenerator()
    gen.render_blocks(blocks)
    assert len(gen.doc.paragraphs) == 3
    texts = [p.text for p in gen.doc.paragraphs]
    assert texts == ["code", "", "Выводы"]

def test_render_code_font():
    """CODE использует Courier New 12."""
    blocks = [Block(kind=BlockKind.CODE, text="code")]
    gen = ReportGenerator()
    gen.render_blocks(blocks)
    run = gen.doc.paragraphs[0].runs[0]
    assert run.font.name == "Courier New"
    assert run.font.size == Pt(12)