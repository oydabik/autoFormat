"""Тесты для list_item"""
import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from generator import ReportGenerator


def test_is_list_item_with_regex():
    """Текст с '1. ' в начале - пункт списка."""
    doc = docx.Document()
    p = doc.add_paragraph("1. Текст")
    gen = ReportGenerator()
    assert gen._is_list_item(p) is True


def test_is_list_item_with_multidigit():
    """Текст с '10. ' в начале - пункт списка."""
    doc = docx.Document()
    p = doc.add_paragraph("10. Текст")
    gen = ReportGenerator()
    assert gen._is_list_item(p) is True


def test_is_list_item_with_plain_text():
    """Текст с 'Текст ' в начале - не пункт списка."""
    doc = docx.Document()
    p = doc.add_paragraph("Текст")
    gen = ReportGenerator()
    assert gen._is_list_item(p) is False


def test_is_list_item_with_text_number():
    """Текст с 'Текст 1. ' в начале - не пункт списка."""
    doc = docx.Document()
    p = doc.add_paragraph("Текст 1. ")
    gen = ReportGenerator()
    assert gen._is_list_item(p) is False


def test_is_list_item_with_numpr():
    """Текст с автонумерацией - пункт списка"""
    doc = docx.Document()
    p = doc.add_paragraph("Текст")
    pPr = p._p.get_or_add_pPr()
    numPr = OxmlElement("w:numPr")
    numId = OxmlElement("w:numId")
    numId.set(qn("w:val"), "1")
    numPr.append(numId)
    pPr.append(numPr)
    gen = ReportGenerator()
    assert gen._is_list_item(p) is True