import pytest
import docx
from docx.shared import Cm


@pytest.fixture
def simple_doc(tmp_path):
    """Простой .docx с несколькими абзацами разных типов."""
    doc = docx.Document()
    doc.add_paragraph("Практика")
    doc.add_paragraph("Задача 1. Текст")
    doc.add_paragraph("Решение.")
    doc.add_paragraph("1. Текст")
    doc.add_paragraph("Текст")

    file_path = tmp_path / "test.docx"
    doc.save(file_path)

    return doc, str(file_path)


@pytest.fixture
def doc_with_15_images(tmp_path):
    """Документ с 15 картинками (номера 1..15 на каждой)."""
    from PIL import Image, ImageDraw, ImageFont

    doc = docx.Document()
    doc.add_paragraph("Тест")

    for i in range(1, 16):
        img = Image.new("RGB", (200, 100), "white")
        draw = ImageDraw.Draw(img)
        try:
            font = ImageFont.truetype(
                "/usr/share/fonts/TTF/DejaVuSans-Bold.ttf", 60
            )
        except OSError:
            font = ImageFont.load_default()
        draw.text((80, 20), str(i), fill="black", font=font)

        img_path = tmp_path / f"test_img_{i}.png"
        img.save(img_path)

        doc.add_paragraph().add_run().add_picture(str(img_path), width=Cm(3))

    file_path = tmp_path / "test.docx"
    doc.save(file_path)
    return str(file_path)