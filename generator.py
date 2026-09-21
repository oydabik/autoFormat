import io
import logging
import xml.etree.ElementTree as ET
import zipfile

import docx
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from config import LIST_ITEM_PATTERN, ReportConfig
from models import Block, BlockKind
from spacing import SpacingRules


logger = logging.getLogger(__name__)
REL_NS = "{http://schemas.openxmlformats.org/package/2006/relationships}Relationship"
EMBED_ATTR = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"


class ReportGenerator:
    def __init__(self):
        self.doc = docx.Document()

    def read_blocks(self, input_doc: docx.Document, docx_path: str) -> list[Block]:
        blocks = []
        is_header = True
        current_mode = "main"
        images = self._extract_images(docx_path)
        image_index = 0

        list_counter = 0
        prev_kind = None

        for paragraph in input_doc.paragraphs:
            text = paragraph.text
            if 'w:drawing' in paragraph._p.xml:
                if image_index < len(images):
                    blocks.append(Block(BlockKind.IMAGE, image=images[image_index]))
                    image_index += 1
                prev_kind = BlockKind.IMAGE
                continue
                
            if not text.strip():
                continue

            if self._is_list_item(paragraph):
                kind = BlockKind.LIST_ITEM
            else:
                kind = self._classify(text, is_header, current_mode)

            if kind == BlockKind.SECTION:
                is_header = False
                current_mode = "main"
            elif kind == BlockKind.SUBHEADER and text.startswith("Исходный код"):
                is_header = False
                current_mode = "code"

            if kind == BlockKind.LIST_ITEM:
                if prev_kind == BlockKind.LIST_ITEM:
                    list_counter += 1
                else:
                    list_counter = 1
                text = f"{list_counter}. {text}"
            else:
                list_counter = 0

            blocks.append(Block(kind, text=text))
            prev_kind = kind
        logger.debug("Прочитано %d блоков", len(blocks))
        return blocks

    def render_blocks(self, blocks: list[Block]) -> None:
        """Отрендерить список блоков в документ."""
        for i, block in enumerate(blocks):
            if block.kind == BlockKind.IMAGE:
                img_stream = io.BytesIO(block.image)
                p_img = self.doc.add_paragraph()
                p_img.add_run().add_picture(img_stream, width=Cm(16.5))
            else:
                self.add_styled_paragraph(block.text, block.kind)

            next_block = blocks[i + 1] if i + 1 < len(blocks) else None
            if SpacingRules.need_empty_after(block, next_block):
                self.doc.add_paragraph()
        logger.debug("Отрендерено %d блоков", len(blocks))

    def _extract_images(self, docx_path: str) -> list[bytes]:
        extracted_images = []
        try:
            with zipfile.ZipFile(docx_path, 'r') as archive:
                # 1. Порядок rId из document.xml
                doc_xml = archive.read("word/document.xml").decode("utf-8")
                doc_root = ET.fromstring(doc_xml)
                rids = [
                    el.get(EMBED_ATTR)
                    for el in doc_root.iter()
                    if el.get(EMBED_ATTR) is not None
                ]

                # 2. Mapping rId → media/... из rels
                rels_xml = archive.read("word/_rels/document.xml.rels").decode("utf-8")
                rels_root = ET.fromstring(rels_xml)
                mapping = {
                    rel.get("Id"): rel.get("Target")
                    for rel in rels_root.findall(REL_NS)
                    if rel.get("Target") and rel.get("Target").startswith("media/")
                }

                # 3. Загружаем картинки в порядке rId
                for rid in rids:
                    target = mapping.get(rid)
                    if target is None:
                        continue
                    img_data = archive.read(f"word/{target}")
                    extracted_images.append(img_data)
        except (zipfile.BadZipFile, FileNotFoundError) as e:
            logger.warning("Не удалось извлечь картинки из %s: %s", docx_path, e)
        logger.debug("Извлечено %d картинок из %s", len(extracted_images), docx_path)
        return extracted_images

    def _classify(self, text: str, is_header: bool, current_mode: str) -> BlockKind:
        if text.startswith(("Задача", "Выводы")):
            return BlockKind.SECTION
        if text.startswith(("Решение.", "Скриншоты", "Исходный код")):
            return BlockKind.SUBHEADER
        
        if is_header:
            return BlockKind.TITLE
        if current_mode == "code":
            return BlockKind.CODE

        if LIST_ITEM_PATTERN.match(text):
            return BlockKind.LIST_ITEM
        
        return BlockKind.BODY

    def _is_list_item(self, paragraph: docx.text.paragraph.Paragraph) -> bool:
        """Проверить, является ли абзац пунктом списка."""

        if paragraph._p.pPr is not None and paragraph._p.pPr.numPr is not None:
            return True
        return LIST_ITEM_PATTERN.match(paragraph.text) is not None

    def add_styled_paragraph(self, text: str, kind:BlockKind) -> None:
        """Добавить абзац с заданным стилем."""
        style = ReportConfig.STYLES[kind]
        p = self.doc.add_paragraph()
        p.alignment = style.alignment
        p.paragraph_format.space_before = Pt(style.space_before)
        p.paragraph_format.space_after = Pt(style.space_after)
        p.paragraph_format.line_spacing = style.line_spacing
        p.paragraph_format.left_indent = Cm(style.left_indent)
        p.paragraph_format.first_line_indent = Cm(style.first_line_indent)

        run = p.add_run(text)
        run.font.name = style.font_name
        rFonts = run._element.get_or_add_rPr().get_or_add_rFonts()
        rFonts.set(qn('w:eastAsia'), style.font_name)
        run.font.size = Pt(style.font_size)

    def parse(self, input_doc: docx.Document, docx_path: str) -> None:
        blocks = self.read_blocks(input_doc, docx_path)
        self.render_blocks(blocks)

    def save(self, file_path: str) -> None:
        self.doc.save(file_path)