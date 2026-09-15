import io
import logging
import re
import zipfile
from dataclasses import dataclass
from enum import Enum, auto

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


logger = logging.getLogger(__name__)

@dataclass
class ParagraphStyle:
    """Параметры форматирования одного абзаца."""
    alignment: WD_ALIGN_PARAGRAPH
    font_name: str
    font_size: int
    line_spacing: float
    space_before: float = 0
    space_after: float = 0
    left_indent: float = 0
    first_line_indent: float = 0


class BlockKind(Enum):
    TITLE = auto()
    SECTION = auto()
    SUBHEADER = auto()
    BODY = auto()
    CODE = auto()
    LIST_ITEM = auto()
    IMAGE = auto()


@dataclass
class Block:
    """Один блок документа (абзац, код, картинка)."""
    kind: BlockKind
    text: str = ""
    image: bytes | None = None


LIST_ITEM_PATTERN = re.compile(r"^\d+\.\s")


class ReportConfig:
    _TEXT = ParagraphStyle(
        alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
        font_name="Times New Roman",
        font_size=14,
        line_spacing=1.5,
    )
    STYLES = {
        BlockKind.TITLE: ParagraphStyle(
            alignment=WD_ALIGN_PARAGRAPH.CENTER,
            font_name="Times New Roman",
            font_size=14,
            line_spacing=1.5,
        ),
        BlockKind.SECTION: _TEXT,
        BlockKind.SUBHEADER: _TEXT,
        BlockKind.BODY: _TEXT,
        BlockKind.CODE: ParagraphStyle(
            alignment=WD_ALIGN_PARAGRAPH.LEFT,
            font_name="Courier New",
            font_size=12,
            line_spacing=1.0,
        ),
        BlockKind.LIST_ITEM: ParagraphStyle(
            alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
            font_name="Times New Roman",
            font_size=14,
            line_spacing=1.5,
            left_indent=1.25,
            first_line_indent=-1.25,
        ),
    }


class SpacingRules:
    """Правила пустых строк между блоками."""
    @staticmethod
    def need_empty_after(prev_block: Block, next_block: Block | None) -> bool:
        if next_block is None:
            return False

        if prev_block.kind == BlockKind.IMAGE and next_block.kind != BlockKind.IMAGE:
            return True

        if prev_block.kind == BlockKind.CODE and next_block.kind != BlockKind.CODE:
            return True

        if prev_block.kind == BlockKind.LIST_ITEM and next_block.kind != BlockKind.LIST_ITEM:
            return True

        return False


class ReportGenerator:
    def __init__(self):
        self.doc = docx.Document()

    def read_blocks(self, input_doc: docx.Document, docx_path: str) -> list[Block]:
        blocks = []
        is_header = True
        current_mode = "main"
        images = self._extract_images(docx_path)
        image_index = 0
        for paragraph in input_doc.paragraphs:
            text = paragraph.text
            if 'w:drawing' in paragraph._p.xml:
                if image_index < len(images):
                    blocks.append(Block(BlockKind.IMAGE, image=images[image_index]))
                    image_index += 1
                continue
                
            if not text.strip():
                continue

            kind = self._classify(text, is_header, current_mode)

            if kind == BlockKind.SECTION:
                is_header = False
                current_mode = "main"
            elif kind == BlockKind.SUBHEADER and text.startswith("Исходный код"):
                is_header = False
                current_mode = "code"

            blocks.append(Block(kind, text=text))
        logger.debug("Прочитано %d блоков", len(blocks))
        return blocks

    def render_blocks(self, blocks: list[Block]) -> None:
        """Отрендерить список блоков в документ."""
        for i, block in enumerate(blocks):
            if block.kind == BlockKind.IMAGE:
                img_stream = io.BytesIO(block.image)
                p_img = self.doc.add_paragraph()
                p_img.add_run().add_picture(img_stream, width=Cm(16.5))
            else:
                self.add_styled_paragraph(block.text, block.kind)

            next_block = blocks[i + 1] if i + 1 < len(blocks) else None
            if SpacingRules.need_empty_after(block, next_block):
                self.doc.add_paragraph()
        logger.debug("Отрендерено %d блоков", len(blocks))

    def _extract_images(self, docx_path: str) -> list[bytes]:
        extracted_images = []
        try:
            with zipfile.ZipFile(docx_path, 'r') as archive:
                media_files = [f for f in archive.namelist() if f.startswith('word/media/')]
                media_files.sort()
                for file_name in media_files:
                    img_data = archive.read(file_name)
                    extracted_images.append(img_data)
        except (zipfile.BadZipFile, FileNotFoundError) as e:
            logger.warning("Не удалось извлечь картинки из %s: %s", docx_path, e)
        logger.debug("Извлечено %d картинок из %s", len(extracted_images), docx_path)
        return extracted_images

    def _classify(self, text: str, is_header: bool, current_mode: str) -> BlockKind:
        if text.startswith(("Задача", "Выводы")):
            return BlockKind.SECTION
        if text.startswith(("Решение.", "Скриншоты", "Исходный код")):
            return BlockKind.SUBHEADER
        
        if is_header:
            return BlockKind.TITLE
        if current_mode == "code":
            return BlockKind.CODE

        if LIST_ITEM_PATTERN.match(text):
            return BlockKind.LIST_ITEM
        
        return BlockKind.BODY

    def add_styled_paragraph(self, text: str, kind:BlockKind) -> None:
        """Добавить абзац с заданным стилем."""
        style = ReportConfig.STYLES[kind]
        p = self.doc.add_paragraph()
        p.alignment = style.alignment
        p.paragraph_format.space_before = Pt(style.space_before)
        p.paragraph_format.space_after = Pt(style.space_after)
        p.paragraph_format.line_spacing = style.line_spacing
        p.paragraph_format.left_indent = Cm(style.left_indent)
        p.paragraph_format.first_line_indent = Cm(style.first_line_indent)

        run = p.add_run(text)
        run.font.name = style.font_name
        rFonts = run._element.get_or_add_rPr().get_or_add_rFonts()
        rFonts.set(qn('w:eastAsia'), style.font_name)
        run.font.size = Pt(style.font_size)

    def parse(self, input_doc: docx.Document, docx_path: str) -> None:
        blocks = self.read_blocks(input_doc, docx_path)
        self.render_blocks(blocks)

    def save(self, file_path: str) -> None:
        self.doc.save(file_path)